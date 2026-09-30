from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
ASSET_DIR = PROJECT_ROOT / "assets"
LOGO_PATH = ASSET_DIR / "logo.png"


# ============================================================
# FILE NAME
# ============================================================

def safe_filename(name: str) -> str:
    name = str(name or "Legal Document")

    name = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        name.strip(),
    )

    name = name.strip("_")

    return name or "legal_document"


def filename_for(
    document_type: str,
    extension: str,
) -> str:
    return (
        f"{safe_filename(document_type)}"
        f".{extension}"
    )


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_pdf_text(text: str) -> str:
    """
    Convert arbitrary document text into text that can be
    safely rendered using FPDF's built-in Helvetica font.
    """

    if text is None:
        return ""

    text = str(text)

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u00a0": " ",
        "\u2022": "-",
        "\u2026": "...",
        "\u00ae": "(R)",
        "\u00a9": "(C)",
        "\u2122": "(TM)",
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u00ad": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Replace any remaining unsupported Unicode characters.
    text = (
        text
        .encode("latin-1", errors="replace")
        .decode("latin-1")
    )

    # Remove control characters.
    cleaned = []

    for char in text:

        if char in ("\n", "\r", "\t"):
            cleaned.append(char)

        elif ord(char) >= 32:
            cleaned.append(char)

    text = "".join(cleaned)

    # Convert tabs to spaces.
    text = text.replace("\t", "    ")

    return text


# ============================================================
# SAFE PDF WRAPPING
# ============================================================

def wrap_text_for_pdf(
    pdf: FPDF,
    text: str,
    max_width: float,
) -> list[str]:
    """
    Wrap text based on actual FPDF string width.

    This does NOT rely on multi_cell() word wrapping,
    which avoids the horizontal-space exception.
    """

    text = clean_pdf_text(text)

    if not text:
        return [""]

    words = text.split(" ")

    lines = []

    current = ""

    for word in words:

        # ----------------------------------------------------
        # Normal word
        # ----------------------------------------------------

        candidate = (
            word
            if not current
            else current + " " + word
        )

        if (
            pdf.get_string_width(candidate)
            <= max_width
        ):

            current = candidate

            continue

        # ----------------------------------------------------
        # Current line is full
        # ----------------------------------------------------

        if current:
            lines.append(current)
            current = ""

        # ----------------------------------------------------
        # Word itself may be too wide.
        # Break it character-by-character.
        # ----------------------------------------------------

        if (
            pdf.get_string_width(word)
            <= max_width
        ):

            current = word

            continue

        chunk = ""

        for character in word:

            candidate_chunk = (
                chunk + character
            )

            if (
                pdf.get_string_width(
                    candidate_chunk
                )
                <= max_width
            ):

                chunk = candidate_chunk

            else:

                if chunk:
                    lines.append(chunk)

                chunk = character

        if chunk:
            current = chunk

    if current:
        lines.append(current)

    return lines


# ============================================================
# TXT
# ============================================================

def make_txt(
    content: str,
) -> BytesIO:

    buffer = BytesIO()

    buffer.write(
        str(content or "").encode(
            "utf-8",
            errors="replace",
        )
    )

    buffer.seek(0)

    return buffer


# ============================================================
# DOCX
# ============================================================

def make_docx(
    document_type: str,
    content: str,
    terms: list[str],
) -> BytesIO:

    document = Document()

    # --------------------------------------------------------
    # Logo
    # --------------------------------------------------------

    if LOGO_PATH.exists():

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                str(LOGO_PATH),
                width=Inches(1.8),
            )

        except Exception:
            pass

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        str(
            document_type
            or "Legal Document"
        )
    )

    title_run.bold = True
    title_run.font.size = Pt(18)

    # --------------------------------------------------------
    # Subtitle
    # --------------------------------------------------------

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle_run = subtitle.add_run(
        "LegalEase - AI-assisted legal document draft"
    )

    subtitle_run.italic = True

    document.add_paragraph()

    # --------------------------------------------------------
    # Main document
    # --------------------------------------------------------

    for raw_line in str(
        content or ""
    ).splitlines():

        line = raw_line.strip()

        if not line:

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph(
            line
        )

        if (
            line.isupper()
            and len(line) < 100
        ):

            for run in paragraph.runs:
                run.bold = True

    # --------------------------------------------------------
    # Terms table
    # --------------------------------------------------------

    if terms:

        document.add_page_break()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "Key Terms"
        )

        heading_run.bold = True
        heading_run.font.size = Pt(14)

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, term in enumerate(
            terms,
            start=1,
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = str(term)

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    for section in document.sections:

        section.footer.paragraphs[0].text = (
            "Generated with LegalEase - "
            "Review before use."
        )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# ============================================================
# PDF CLASS
# ============================================================

class LegalEasePDF(FPDF):

    def __init__(
        self,
        document_type: str,
    ):

        super().__init__()

        self.document_type = clean_pdf_text(
            document_type
            or "Legal Document"
        )

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=10,
                    y=8,
                    w=25,
                )

            except Exception:
                pass

        self.set_font(
            "Helvetica",
            "B",
            15,
        )

        title = self.document_type

        # Use a safe width.
        available_width = 145

        title_lines = wrap_text_for_pdf(
            self,
            title,
            available_width,
        )

        self.set_x(45)

        if title_lines:

            self.cell(
                available_width,
                8,
                title_lines[0][:80],
                align="C",
            )

        self.ln(15)

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8,
        )

        self.cell(
            0,
            10,
            "LegalEase - Review before use.",
            align="C",
        )


# ============================================================
# PDF
# ============================================================

def make_pdf(
    document_type: str,
    content: str,
) -> BytesIO:

    pdf = LegalEasePDF(
        document_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "",
        11,
    )

    # --------------------------------------------------------
    # Available body width
    # --------------------------------------------------------

    left_margin = 10
    right_margin = 10

    usable_width = (
        pdf.w
        - left_margin
        - right_margin
    )

    # Keep a small safety margin.
    max_width = usable_width - 4

    # --------------------------------------------------------
    # Render document line by line
    # --------------------------------------------------------

    raw_lines = str(
        content or ""
    ).replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    ).split("\n")

    for raw_line in raw_lines:

        line = raw_line.strip()

        # Empty line
        if not line:

            pdf.ln(5)

            continue

        # ----------------------------------------------------
        # Wrap using actual measured character widths.
        # ----------------------------------------------------

        wrapped_lines = wrap_text_for_pdf(
            pdf,
            line,
            max_width,
        )

        for wrapped_line in wrapped_lines:

            if not wrapped_line:
                continue

            # ------------------------------------------------
            # NEVER use multi_cell here.
            # ------------------------------------------------

            pdf.cell(
                max_width,
                7,
                wrapped_line,
                new_x="LMARGIN",
                new_y="NEXT",
            )

    # --------------------------------------------------------
    # Output
    # --------------------------------------------------------

    data = bytes(
        pdf.output()
    )

    buffer = BytesIO(
        data
    )

    buffer.seek(0)

    return buffer